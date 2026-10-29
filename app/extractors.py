import re
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from .ai.ollama import extract_jd as ollama_extract_jd
from .ai.ollama import extract_resume as ollama_extract_resume
from .models import CandidateProfile, ExtractedDocument, Skill

CATEGORIES: dict[str, tuple[str, ...]] = {
    "COD": ("python", "java", "javascript", "typescript", "c++", "c#", "go", "ruby", "programming"),
    "DSA": ("data structures", "algorithms", "dsa"),
    "OOD": ("object-oriented", "oop", "design patterns"),
    "APTI": ("aptitude", "problem solving", "analytical"),
    "COMM": ("communication", "collaboration", "leadership", "stakeholder"),
    "AI": ("machine learning", "deep learning", "artificial intelligence", "nlp", "tensorflow", "pytorch"),
    "CLOUD": ("aws", "azure", "gcp", "google cloud", "docker", "kubernetes", "cloud"),
    "SQL": ("sql", "mysql", "postgresql", "oracle database", "mongodb", "database"),
    "SWE": ("software engineering", "rest api", "api", "git", "testing", "agile", "spring", "react"),
    "SYSD": ("system design", "distributed systems", "microservices", "scalability"),
    "NETW": ("networking", "tcp/ip", "http", "dns"),
    "OS": ("linux", "unix", "operating systems", "windows"),
}

SECTION_NAMES = {
    "education", "experience", "work experience", "professional experience",
    "projects", "technical skills", "skills", "certifications", "internships",
    "hackathons", "achievements", "summary", "objective",
}


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        text = "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
    elif suffix == ".docx":
        document = Document(str(path))
        text = "\n".join([p.text for p in document.paragraphs] + [
            " | ".join(cell.text for cell in row.cells)
            for table in document.tables for row in table.rows
        ])
    else:
        raise ValueError("Unsupported document format. Upload a PDF or DOCX file.")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not text:
        raise ValueError("The document contains no extractable text. Scanned files need OCR before upload.")
    return text


def _skill_list(text: str) -> list[Skill]:
    lowered = text.lower()
    result: list[Skill] = []
    seen: set[str] = set()
    for category, terms in CATEGORIES.items():
        for term in terms:
            if term in lowered and term not in seen:
                match = re.search(re.escape(term), text, re.IGNORECASE)
                evidence = match.group(0) if match else term
                result.append(Skill(
                    skill_name=evidence,
                    category_code=category, evidence=f"Explicitly stated: {evidence}",
                    confidence="high",
                ))
                seen.add(term)
    return result


def _metadata_from_filename(filename: str) -> tuple[str | None, str | None]:
    stem = Path(filename).stem
    for company in ("Google LLC", "Microsoft", "Oracle"):
        if stem.lower().startswith(company.lower()):
            return company, stem[len(company):].strip(" -")
    return None, stem


def parse_jd(filename: str, text: str, document_id: str) -> ExtractedDocument:
    extracted = ollama_extract_jd(text)
    company = extracted.company
    role = extracted.role
    filename_company, filename_role = _metadata_from_filename(filename)
    company = company or filename_company
    role = role or filename_role
    warnings: list[str] = []
    missing: list[str] = []
    if not company:
        missing.append("company")
        warnings.append("Company could not be identified from the filename.")
    if not role:
        missing.append("role")
    skills = extracted.skills
    if not skills:
        warnings.append("The local model did not identify any skills in the document.")
    return ExtractedDocument(
        id=document_id, source_type="jd", source_file=filename, company=company,
        role=role, skills=skills, raw_text=text, missing_fields=missing, warnings=warnings,
        extraction_method="ollama",
    )


def _sections(text: str) -> dict[str, list[str]]:
    current = "summary"
    sections: dict[str, list[str]] = {current: []}
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        normalized = re.sub(r"[^a-z ]", "", line.lower()).strip()
        if normalized in SECTION_NAMES:
            current = normalized
            sections.setdefault(current, [])
        else:
            sections.setdefault(current, []).append(line)
    return sections


def _first_email(text: str) -> str | None:
    match = re.search(r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b", text, re.I)
    return match.group(0) if match else None


def _first_phone(text: str) -> str | None:
    match = re.search(r"(?<!\d)(?:\+?91[\s-]?)?[6-9]\d{9}(?!\d)", re.sub(r"[()]", "", text))
    return match.group(0) if match else None


def parse_resume(filename: str, text: str, document_id: str) -> tuple[ExtractedDocument, CandidateProfile]:
    parsed = ollama_extract_resume(text)
    name = parsed.name or Path(filename).stem
    skills = parsed.skills
    missing: list[str] = []
    warnings: list[str] = []
    email = parsed.email
    phone = parsed.phone
    if not email:
        missing.append("email")
    education = parsed.education
    experience = parsed.experience
    projects = parsed.projects
    certifications = parsed.certifications
    internships = parsed.internships
    hackathons = parsed.hackathons
    if not education:
        missing.append("education")
    if not skills:
        warnings.append("No supported skill terms were explicitly detected.")
    extracted = ExtractedDocument(
        id=document_id, source_type="resume", source_file=filename, skills=skills,
        raw_text=text, missing_fields=missing, warnings=warnings, extraction_method="ollama",
    )
    profile = CandidateProfile(
        id=f"profile-{document_id}", resume_id=document_id, name=name, email=email, phone=phone,
        education=education, skills=skills, experience=experience, projects=projects,
        hackathons=hackathons, internships=internships, certifications=certifications,
        preferred_roles=parsed.preferred_roles, cv_file=filename, missing_fields=missing, warnings=warnings,
    )
    return extracted, profile
